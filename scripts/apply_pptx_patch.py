#!/usr/bin/env python3
"""Apply deterministic patches with precondition checks to PowerPoint (.pptx) files.

Supported ops:
  - replace_text: Find & replace in shapes, tables, or notes with expected_matches validation.
  - set_slide_title: Update slide title with optional expected_old_title validation.
  - set_slide_bullets: Replace bullet points in content placeholder while preserving font style.
  - set_speaker_notes: Create or replace/append speaker notes with optional expected_contains.
  - fill_table: Cell-targeted table filling using 1-based row/col coordinates.
  - replace_image: Replace image blob in existing picture shapes or placeholders.

Usage:
  python scripts/apply_pptx_patch.py --in presentation.pptx --out patched.pptx --patch patch.json
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER


def existing_pptx(path_str: str) -> Path:
    p = Path(path_str)
    if not p.exists():
        raise argparse.ArgumentTypeError(f"Datei nicht gefunden: {path_str}")
    if p.suffix.lower() != ".pptx":
        raise argparse.ArgumentTypeError(f"Keine .pptx-Datei: {path_str}")
    return p


def existing_json(path_str: str) -> Path:
    p = Path(path_str)
    if not p.exists():
        raise argparse.ArgumentTypeError(f"Patch-Datei nicht gefunden: {path_str}")
    if p.suffix.lower() not in {".json", ".jsonl"}:
        raise argparse.ArgumentTypeError(f"Keine JSON-Datei: {path_str}")
    return p


def get_slide_by_number(prs: Presentation, slide_number: int) -> Any:
    if slide_number < 1 or slide_number > len(prs.slides):
        raise ValueError(f"Folie {slide_number} existiert nicht (Gesamt: {len(prs.slides)}).")
    return prs.slides[slide_number - 1]


def find_shape_by_id(slide: Any, shape_id: int) -> Any:
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    raise ValueError(f"Shape mit ID {shape_id} auf Folie nicht gefunden.")


def find_body_placeholder(slide: Any) -> Any:
    for shape in slide.placeholders:
        try:
            pt = shape.placeholder_format.type
            if pt in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
                return shape
        except Exception:
            pass
    for shape in slide.shapes:
        if shape.has_text_frame and shape != slide.shapes.title:
            return shape
    raise ValueError("Kein Inhalts-Platzhalter oder Textfeld auf der Folie gefunden.")


def find_table_shape(slide: Any, shape_id: int | None = None) -> Any:
    if shape_id is not None:
        shape = find_shape_by_id(slide, shape_id)
        if not shape.has_table:
            raise ValueError(f"Shape {shape_id} ist keine Tabelle.")
        return shape
    for shape in slide.shapes:
        if shape.has_table:
            return shape
    raise ValueError("Keine Tabelle auf der Folie gefunden.")


def find_picture_shape(slide: Any, shape_id: int | None = None) -> Any:
    if shape_id is not None:
        return find_shape_by_id(slide, shape_id)
    for shape in slide.placeholders:
        try:
            if shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
                return shape
        except Exception:
            pass
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return shape
    raise ValueError("Kein Bild oder Bild-Platzhalter auf der Folie gefunden.")


def op_replace_text(prs: Presentation, op: dict[str, Any]) -> dict[str, Any]:
    find_str = op.get("find")
    replace_str = op.get("replace", "")
    expected_matches = int(op.get("expected_matches", 1))
    target_slide = op.get("slide_number")
    target_shape_id = op.get("shape_id")
    include_notes = bool(op.get("include_notes", False))

    if not isinstance(find_str, str) or not find_str:
        raise ValueError("replace_text erfordert nicht-leeres 'find'.")
    if not isinstance(replace_str, str):
        raise ValueError("replace_text erfordert String in 'replace'.")

    total_matches = 0
    slides_to_scan = []
    if target_slide is not None:
        slides_to_scan.append((target_slide, get_slide_by_number(prs, target_slide)))
    else:
        slides_to_scan.extend(enumerate(prs.slides, start=1))

    for s_idx, slide in slides_to_scan:
        shapes_to_scan = []
        if target_shape_id is not None:
            shapes_to_scan.append(find_shape_by_id(slide, target_shape_id))
        else:
            shapes_to_scan.extend(slide.shapes)

        for shape in shapes_to_scan:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        if find_str in run.text:
                            c = run.text.count(find_str)
                            run.text = run.text.replace(find_str, replace_str)
                            total_matches += c
                    if find_str in p.text and not any(find_str in r.text for r in p.runs):
                        c = p.text.count(find_str)
                        p.text = p.text.replace(find_str, replace_str)
                        total_matches += c

            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                if find_str in run.text:
                                    c = run.text.count(find_str)
                                    run.text = run.text.replace(find_str, replace_str)
                                    total_matches += c
                            if find_str in p.text and not any(find_str in r.text for r in p.runs):
                                c = p.text.count(find_str)
                                p.text = p.text.replace(find_str, replace_str)
                                total_matches += c

        if include_notes and slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            for p in notes_tf.paragraphs:
                for run in p.runs:
                    if find_str in run.text:
                        c = run.text.count(find_str)
                        run.text = run.text.replace(find_str, replace_str)
                        total_matches += c
                if find_str in p.text and not any(find_str in r.text for r in p.runs):
                    c = p.text.count(find_str)
                    p.text = p.text.replace(find_str, replace_str)
                    total_matches += c

    if total_matches != expected_matches:
        raise ValueError(
            f"replace_text: expected_matches={expected_matches}, aber {total_matches} Treffer gefunden. "
            "Abbruch zur Vermeidung inkonsistenter Änderungen."
        )

    return {
        "op": "replace_text",
        "status": "ok",
        "matches": total_matches,
        "find": find_str,
        "replace": replace_str,
    }


def op_set_slide_title(prs: Presentation, op: dict[str, Any]) -> dict[str, Any]:
    slide_number = int(op["slide_number"])
    new_title = str(op["title"])
    expected_old = op.get("expected_old_title")

    slide = get_slide_by_number(prs, slide_number)
    title_shape = slide.shapes.title

    if title_shape is None:
        raise ValueError(f"Folie {slide_number} hat keinen Titel-Platzhalter.")

    current_title = title_shape.text.strip()
    if expected_old is not None:
        if expected_old not in current_title:
            raise ValueError(
                f"set_slide_title auf Folie {slide_number}: expected_old_title '{expected_old}' "
                f"nicht in aktuellem Titel '{current_title}' enthalten."
            )

    title_shape.text = new_title
    return {
        "op": "set_slide_title",
        "status": "ok",
        "slide_number": slide_number,
        "old_title": current_title,
        "new_title": new_title,
    }


def op_set_slide_bullets(prs: Presentation, op: dict[str, Any]) -> dict[str, Any]:
    slide_number = int(op["slide_number"])
    bullets = op.get("bullets")
    shape_id = op.get("shape_id")
    expected_contains = op.get("expected_contains")

    if not isinstance(bullets, list):
        raise ValueError("set_slide_bullets erfordert eine Liste in 'bullets'.")

    slide = get_slide_by_number(prs, slide_number)
    shape = find_shape_by_id(slide, shape_id) if shape_id is not None else find_body_placeholder(slide)

    if not shape.has_text_frame:
        raise ValueError(f"Shape auf Folie {slide_number} hat keinen Text-Frame.")

    tf = shape.text_frame
    current_text = tf.text.strip()

    if expected_contains is not None and expected_contains not in current_text:
        raise ValueError(
            f"set_slide_bullets auf Folie {slide_number}: expected_contains '{expected_contains}' "
            "nicht im bestehenden Text gefunden."
        )

    # Capture font styling from original paragraph if available
    font_name = None
    font_size = None
    font_color_rgb = None
    if tf.paragraphs and tf.paragraphs[0].runs:
        r0 = tf.paragraphs[0].runs[0]
        font_name = r0.font.name
        font_size = r0.font.size
        try:
            if r0.font.color and r0.font.color.rgb:
                font_color_rgb = r0.font.color.rgb
        except Exception:
            pass

    # Clear and rewrite
    tf.clear()
    for idx, b_item in enumerate(bullets):
        text = ""
        level = 0
        bold = False
        italic = False

        if isinstance(b_item, str):
            text = b_item
        elif isinstance(b_item, dict):
            text = str(b_item.get("text", ""))
            level = int(b_item.get("level", 0))
            bold = bool(b_item.get("bold", False))
            italic = bool(b_item.get("italic", False))
        else:
            raise ValueError(f"Ungültiger Bullet-Eintrag: {b_item}")

        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.level = level
        if p.runs:
            r = p.runs[0]
            if bold:
                r.font.bold = True
            if italic:
                r.font.italic = True
            if font_name:
                r.font.name = font_name
            if font_size:
                r.font.size = font_size
            if font_color_rgb:
                try:
                    r.font.color.rgb = font_color_rgb
                except Exception:
                    pass

    return {
        "op": "set_slide_bullets",
        "status": "ok",
        "slide_number": slide_number,
        "bullets_count": len(bullets),
    }


def op_set_speaker_notes(prs: Presentation, op: dict[str, Any]) -> dict[str, Any]:
    slide_number = int(op["slide_number"])
    new_text = str(op["text"])
    mode = str(op.get("mode", "replace"))
    expected_contains = op.get("expected_contains")

    slide = get_slide_by_number(prs, slide_number)
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    current_notes = tf.text.strip()

    if expected_contains is not None and expected_contains not in current_notes:
        raise ValueError(
            f"set_speaker_notes auf Folie {slide_number}: expected_contains '{expected_contains}' "
            "nicht in aktuellen Notizen gefunden."
        )

    if mode == "append" and current_notes:
        tf.text = f"{current_notes}\n{new_text}"
    else:
        tf.text = new_text

    return {
        "op": "set_speaker_notes",
        "status": "ok",
        "slide_number": slide_number,
        "mode": mode,
        "length": len(tf.text),
    }


def op_fill_table(prs: Presentation, op: dict[str, Any]) -> dict[str, Any]:
    slide_number = int(op["slide_number"])
    shape_id = op.get("shape_id")
    cells = op.get("cells")

    if not isinstance(cells, list):
        raise ValueError("fill_table erfordert eine Liste in 'cells'.")

    slide = get_slide_by_number(prs, slide_number)
    shape = find_table_shape(slide, shape_id)
    tbl = shape.table

    rows_count = len(tbl.rows)
    cols_count = len(tbl.columns)
    updated_count = 0

    for c in cells:
        row_1b = int(c["row"])
        col_1b = int(c["col"])
        c_text = str(c.get("text", ""))
        c_mode = str(c.get("mode", "replace"))

        if row_1b < 1 or row_1b > rows_count:
            raise ValueError(f"Zeile {row_1b} außerhalb des Tabellenbereichs (1..{rows_count}).")
        if col_1b < 1 or col_1b > cols_count:
            raise ValueError(f"Spalte {col_1b} außerhalb des Tabellenbereichs (1..{cols_count}).")

        cell = tbl.cell(row_1b - 1, col_1b - 1)
        if c_mode == "append" and cell.text:
            cell.text = f"{cell.text}\n{c_text}"
        else:
            cell.text = c_text
        updated_count += 1

    return {
        "op": "fill_table",
        "status": "ok",
        "slide_number": slide_number,
        "cells_updated": updated_count,
    }


def op_replace_image(prs: Presentation, op: dict[str, Any]) -> dict[str, Any]:
    slide_number = int(op["slide_number"])
    shape_id = op.get("shape_id")
    image_path_str = str(op["image_path"])

    img_path = Path(image_path_str)
    if not img_path.exists():
        raise ValueError(f"Bilddatei nicht gefunden: {image_path_str}")

    img_bytes = img_path.read_bytes()
    slide = get_slide_by_number(prs, slide_number)
    shape = find_picture_shape(slide, shape_id)

    # 1. If it is a placeholder that is unfilled
    if shape.is_placeholder:
        try:
            shape.insert_picture(str(img_path))
            return {
                "op": "replace_image",
                "status": "ok",
                "slide_number": slide_number,
                "action": "inserted_into_placeholder",
            }
        except Exception:
            pass

    # 2. If it is an existing picture shape, swap underlying ImagePart blob
    rel_ids = shape._element.xpath(".//a:blip/@r:embed")
    if not rel_ids:
        raise ValueError(f"Shape auf Folie {slide_number} enthält keine Bild-Referenz (a:blip).")

    rel_id = rel_ids[0]
    img_part = slide.part.related_part(rel_id)
    img_part._blob = img_bytes

    return {
        "op": "replace_image",
        "status": "ok",
        "slide_number": slide_number,
        "action": "blob_replaced",
        "image_bytes": len(img_bytes),
    }


def apply_patch(infile: Path, outfile: Path, patch_file: Path) -> dict[str, Any]:
    patch_data = json.loads(patch_file.read_text(encoding="utf-8"))
    ops = patch_data.get("ops")
    if not isinstance(ops, list):
        raise ValueError("Patch-Datei muss ein Root-Objekt mit Array 'ops' enthalten.")

    prs = Presentation(str(infile))
    results: list[dict[str, Any]] = []

    for idx, op in enumerate(ops, start=1):
        if not isinstance(op, dict):
            raise ValueError(f"Op #{idx} ist kein Objekt.")
        op_name = op.get("op")
        if op_name == "replace_text":
            res = op_replace_text(prs, op)
        elif op_name == "set_slide_title":
            res = op_set_slide_title(prs, op)
        elif op_name == "set_slide_bullets":
            res = op_set_slide_bullets(prs, op)
        elif op_name == "set_speaker_notes":
            res = op_set_speaker_notes(prs, op)
        elif op_name == "fill_table":
            res = op_fill_table(prs, op)
        elif op_name == "replace_image":
            res = op_replace_image(prs, op)
        else:
            raise ValueError(f"Unbekannte Operation #{idx}: '{op_name}'")
        results.append(res)

    # Atomic write to outfile
    outfile.parent.mkdir(parents=True, exist_ok=True)
    temp_file = tempfile.NamedTemporaryFile(
        dir=str(outfile.parent),
        prefix=f".tmp_{outfile.stem}_",
        suffix=".pptx",
        delete=False,
    )
    temp_path = Path(temp_file.name)
    temp_file.close()

    try:
        prs.save(str(temp_path))
        os.replace(str(temp_path), str(outfile))
    finally:
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass

    return {
        "action": "apply_patch",
        "success": True,
        "message": f"Applied {len(results)} operation(s) successfully",
        "in": str(infile).replace("\\", "/"),
        "out": str(outfile).replace("\\", "/"),
        "results": results,
    }


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Apply deterministic patches to PPTX")
    p.add_argument("--in", dest="infile", type=existing_pptx, required=True, help="Pfad zur PPTX-Eingabedatei")
    p.add_argument("--out", dest="outfile", type=Path, required=True, help="Pfad zur Ausgabedatei")
    p.add_argument("--patch", dest="patchfile", type=existing_json, required=True, help="Pfad zur Patch-JSON-Datei")
    return p.parse_args()


def main() -> int:
    args = parse_args()
    res = apply_patch(args.infile, args.outfile, args.patchfile)
    print(json.dumps(res, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # noqa: BLE001
        err_res = {
            "action": "apply_patch",
            "success": False,
            "error": str(exc),
        }
        print(json.dumps(err_res, ensure_ascii=False, indent=2), file=sys.stderr)
        raise SystemExit(2)
