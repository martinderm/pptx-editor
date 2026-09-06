#!/usr/bin/env python3
"""Extract a PowerPoint presentation (.pptx) into structured JSON for LLM workflows.

Default output (v1): Hierarchical presentation structure grouped by slides, shapes, tables, and speaker notes.
Optional RAG output: Flat blocks + chunks via --rag-output for vector retrieval.

Usage:
  python scripts/extract_pptx_for_llm.py --in input.pptx --out structure.json
  python scripts/extract_pptx_for_llm.py --in input.pptx --out structure.json --rag-output rag.json
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


@dataclass
class RAGBlock:
    block_id: str
    slide_number: int
    block_type: str
    title: str
    text: str


def existing_pptx(path_str: str) -> Path:
    p = Path(path_str)
    if not p.exists():
        raise argparse.ArgumentTypeError(f"Datei nicht gefunden: {path_str}")
    if p.suffix.lower() != ".pptx":
        raise argparse.ArgumentTypeError(f"Keine .pptx-Datei: {path_str}")
    return p


def get_shape_type_name(shape: Any) -> str:
    if shape.is_placeholder:
        return "placeholder"
    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.TEXT_BOX:
        return "text_box"
    if st == MSO_SHAPE_TYPE.TABLE:
        return "table"
    if st == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
        return "auto_shape"
    if st == MSO_SHAPE_TYPE.GROUP:
        return "group"
    return "other"


def get_placeholder_type_name(shape: Any) -> str | None:
    if not shape.is_placeholder:
        return None
    try:
        pt = shape.placeholder_format.type
        return str(pt).split(".")[-1].split(" ")[0]
    except Exception:
        return "UNKNOWN"


def extract_runs(paragraph: Any) -> list[dict[str, Any]]:
    runs_data: list[dict[str, Any]] = []
    for r in paragraph.runs:
        t = r.text
        if not t:
            continue
        run_info: dict[str, Any] = {
            "text": t,
            "bold": bool(r.font.bold) if r.font.bold is not None else False,
            "italic": bool(r.font.italic) if r.font.italic is not None else False,
        }
        if r.font.name:
            run_info["font"] = r.font.name
        if r.font.size is not None:
            run_info["size_pt"] = round(r.font.size.pt, 1)
        try:
            if r.font.color and r.font.color.rgb:
                run_info["color"] = str(r.font.color.rgb)
        except Exception:
            pass
        runs_data.append(run_info)
    return runs_data


def extract_v1(infile: Path, prs: Presentation) -> dict[str, Any]:
    width_in = prs.slide_width / 914400.0
    height_in = prs.slide_height / 914400.0
    width_pt = prs.slide_width / 12700.0
    height_pt = prs.slide_height / 12700.0

    slides_data: list[dict[str, Any]] = []
    total_words = 0
    total_shapes = 0
    total_tables = 0
    total_images = 0
    empty_placeholders = 0

    for idx, slide in enumerate(prs.slides, start=1):
        slide_title = None
        if slide.shapes.title and slide.shapes.title.text:
            slide_title = slide.shapes.title.text.strip()
            total_words += len(slide_title.split())

        shapes_data: list[dict[str, Any]] = []
        tables_data: list[dict[str, Any]] = []

        for shape in slide.shapes:
            total_shapes += 1
            stype = get_shape_type_name(shape)
            ph_type = get_placeholder_type_name(shape)

            if shape.is_placeholder and shape.has_text_frame and not shape.text_frame.text.strip():
                empty_placeholders += 1

            if stype == "picture":
                total_images += 1

            pos = {
                "left_pt": round(shape.left / 12700.0, 1),
                "top_pt": round(shape.top / 12700.0, 1),
                "width_pt": round(shape.width / 12700.0, 1),
                "height_pt": round(shape.height / 12700.0, 1),
            }

            if shape.has_table:
                total_tables += 1
                tbl = shape.table
                table_cells: list[list[dict[str, Any]]] = []
                for r_idx, row in enumerate(tbl.rows, start=1):
                    row_cells: list[dict[str, Any]] = []
                    for c_idx, cell in enumerate(row.cells, start=1):
                        c_text = cell.text.strip()
                        if c_text:
                            total_words += len(c_text.split())
                        row_cells.append({
                            "row": r_idx,
                            "col": c_idx,
                            "text": c_text,
                        })
                    table_cells.append(row_cells)

                tables_data.append({
                    "table_id": f"s{idx}_tbl{shape.shape_id}",
                    "shape_id": shape.shape_id,
                    "name": shape.name,
                    "position_pt": pos,
                    "rows_count": len(tbl.rows),
                    "cols_count": len(tbl.columns),
                    "cells": table_cells,
                })

            elif shape.has_text_frame:
                paragraphs_data: list[dict[str, Any]] = []
                for p_idx, p in enumerate(shape.text_frame.paragraphs, start=1):
                    p_text = p.text.strip()
                    if p_text:
                        total_words += len(p_text.split())
                    paragraphs_data.append({
                        "p_id": f"s{idx}_sh{shape.shape_id}_p{p_idx}",
                        "level": p.level,
                        "text": p_text,
                        "runs": extract_runs(p),
                    })

                shapes_data.append({
                    "shape_id": shape.shape_id,
                    "shape_name": shape.name,
                    "shape_type": stype,
                    "is_placeholder": shape.is_placeholder,
                    "placeholder_type": ph_type,
                    "position_pt": pos,
                    "text": shape.text_frame.text.strip() if shape.text_frame.text else "",
                    "paragraphs": paragraphs_data,
                })

            elif stype == "picture":
                shapes_data.append({
                    "shape_id": shape.shape_id,
                    "shape_name": shape.name,
                    "shape_type": "picture",
                    "is_placeholder": shape.is_placeholder,
                    "placeholder_type": ph_type,
                    "position_pt": pos,
                })

        # Speaker notes
        notes_dict: dict[str, Any] = {
            "has_notes": False,
            "text": "",
            "paragraphs": [],
        }
        if slide.has_notes_slide:
            ntf = slide.notes_slide.notes_text_frame
            n_text = ntf.text.strip()
            if n_text:
                notes_dict["has_notes"] = True
                notes_dict["text"] = n_text
                total_words += len(n_text.split())
                for np_idx, np in enumerate(ntf.paragraphs, start=1):
                    np_t = np.text.strip()
                    if np_t:
                        notes_dict["paragraphs"].append({
                            "p_id": f"s{idx}_notes_p{np_idx}",
                            "text": np_t,
                        })

        slides_data.append({
            "slide_number": idx,
            "slide_id": slide.slide_id,
            "title": slide_title,
            "layout_name": slide.slide_layout.name,
            "shapes": shapes_data,
            "tables": tables_data,
            "speaker_notes": notes_dict,
        })

    return {
        "schema": "pptx-structure.v1",
        "source": str(infile).replace("\\", "/"),
        "dimensions": {
            "width_pt": round(width_pt, 1),
            "height_pt": round(height_pt, 1),
            "width_inches": round(width_in, 2),
            "height_inches": round(height_in, 2),
        },
        "stats": {
            "total_slides": len(prs.slides),
            "total_words": total_words,
            "shapes_count": total_shapes,
            "tables_count": total_tables,
            "images_count": total_images,
            "empty_placeholders": empty_placeholders,
        },
        "slides": slides_data,
    }


def extract_rag_blocks(v1_data: dict[str, Any]) -> list[RAGBlock]:
    blocks: list[RAGBlock] = []

    for s in v1_data["slides"]:
        s_num = s["slide_number"]
        s_title = s["title"] or f"Slide {s_num}"

        if s["title"]:
            blocks.append(RAGBlock(
                block_id=f"s{s_num}_title",
                slide_number=s_num,
                block_type="slide_title",
                title=s_title,
                text=f"# {s_title}",
            ))

        body_lines: list[str] = []
        for sh in s.get("shapes", []):
            if sh.get("shape_name", "").startswith("Title"):
                continue
            for p in sh.get("paragraphs", []):
                t = p.get("text", "").strip()
                if t:
                    indent = "  " * p.get("level", 0)
                    body_lines.append(f"{indent}- {t}")

        if body_lines:
            blocks.append(RAGBlock(
                block_id=f"s{s_num}_body",
                slide_number=s_num,
                block_type="slide_content",
                title=s_title,
                text="\n".join(body_lines),
            ))

        for tbl in s.get("tables", []):
            tbl_lines: list[str] = []
            for row in tbl.get("cells", []):
                tbl_lines.append(" | ".join(c.get("text", "") for c in row))
            if tbl_lines:
                blocks.append(RAGBlock(
                    block_id=tbl["table_id"],
                    slide_number=s_num,
                    block_type="table",
                    title=f"{s_title} - Table",
                    text="\n".join(tbl_lines),
                ))

        notes = s.get("speaker_notes", {})
        if notes.get("has_notes") and notes.get("text"):
            blocks.append(RAGBlock(
                block_id=f"s{s_num}_notes",
                slide_number=s_num,
                block_type="speaker_notes",
                title=f"{s_title} - Speaker Notes",
                text=notes["text"],
            ))

    return blocks


def chunk_rag_blocks(blocks: list[RAGBlock], max_chars: int = 4000, overlap: int = 1) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    current_blocks: list[RAGBlock] = []
    current_len = 0
    chunk_idx = 1

    i = 0
    while i < len(blocks):
        b = blocks[i]
        b_len = len(b.text)

        if current_blocks and (current_len + b_len > max_chars):
            chunks.append({
                "chunk_id": f"c_{chunk_idx}",
                "slide_numbers": sorted(list({x.slide_number for x in current_blocks})),
                "block_ids": [x.block_id for x in current_blocks],
                "text": "\n\n".join(x.text for x in current_blocks),
            })
            chunk_idx += 1

            if overlap > 0 and len(current_blocks) > overlap:
                current_blocks = current_blocks[-overlap:]
                current_len = sum(len(x.text) for x in current_blocks)
            else:
                current_blocks = []
                current_len = 0

        current_blocks.append(b)
        current_len += b_len
        i += 1

    if current_blocks:
        chunks.append({
            "chunk_id": f"c_{chunk_idx}",
            "slide_numbers": sorted(list({x.slide_number for x in current_blocks})),
            "block_ids": [x.block_id for x in current_blocks],
            "text": "\n\n".join(x.text for x in current_blocks),
        })

    return chunks


def build_rag_json(source: Path, blocks: list[RAGBlock], chunks: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "schema": "pptx-rag-chunks.v1",
        "source": str(source).replace("\\", "/"),
        "stats": {
            "blocks": len(blocks),
            "chunks": len(chunks),
        },
        "blocks": [
            {
                "block_id": b.block_id,
                "slide_number": b.slide_number,
                "type": b.block_type,
                "title": b.title,
                "text": b.text,
            }
            for b in blocks
        ],
        "chunks": chunks,
    }


def write_json_atomic(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_file = tempfile.NamedTemporaryFile(
        dir=str(path.parent),
        prefix=f".tmp_{path.stem}_",
        suffix=".json",
        delete=False,
    )
    temp_path = Path(temp_file.name)
    temp_file.close()

    try:
        temp_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(str(temp_path), str(path))
    finally:
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Extract PPTX as structured JSON for LLM workflows")
    p.add_argument("--in", dest="infile", type=existing_pptx, required=True, help="Pfad zur PPTX-Datei")
    p.add_argument("--out", dest="outfile", type=Path, required=True, help="Ausgabepfad für pptx-structure.v1")
    p.add_argument("--rag-output", dest="rag_outfile", type=Path, help="Optionaler Ausgabepfad für RAG-Chunks")
    p.add_argument("--max-chars", type=int, default=4000, help="Maximale Zeichenanzahl pro RAG-Chunk")
    p.add_argument("--overlap-blocks", type=int, default=1, help="Block-Überlappung für RAG-Chunks")
    return p.parse_args()


def main() -> int:
    args = parse_args()
    prs = Presentation(str(args.infile))

    v1 = extract_v1(args.infile, prs)
    write_json_atomic(args.outfile, v1)

    result: dict[str, Any] = {
        "action": "extract",
        "success": True,
        "out": str(args.outfile).replace("\\", "/"),
        "schema": v1["schema"],
        "slides": v1["stats"]["total_slides"],
        "words": v1["stats"]["total_words"],
    }

    if args.rag_outfile:
        blocks = extract_rag_blocks(v1)
        chunks = chunk_rag_blocks(blocks, max_chars=max(500, args.max_chars), overlap=max(0, args.overlap_blocks))
        rag_data = build_rag_json(args.infile, blocks, chunks)
        write_json_atomic(args.rag_outfile, rag_data)
        result["rag_output"] = str(args.rag_outfile).replace("\\", "/")
        result["rag_blocks"] = rag_data["stats"]["blocks"]
        result["rag_chunks"] = rag_data["stats"]["chunks"]

    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # noqa: BLE001
        err_res = {
            "action": "extract",
            "success": False,
            "error": str(exc),
        }
        print(json.dumps(err_res, ensure_ascii=False, indent=2), file=sys.stderr)
        raise SystemExit(2)
