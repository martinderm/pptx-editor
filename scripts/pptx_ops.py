#!/usr/bin/env python3
"""Basic CLI operations for PowerPoint (.pptx) presentations using python-pptx.

Usage:
  python scripts/pptx_ops.py text --in presentation.pptx [--no-notes] [--json]
  python scripts/pptx_ops.py stats --in presentation.pptx [--json]
  python scripts/pptx_ops.py replace --in in.pptx --out out.pptx --find "Old" --replace "New" [--slide 1] [--include-notes] [--json]
  python scripts/pptx_ops.py inspect --in presentation.pptx [--slide 1] [--json]
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Inches, Pt


def make_envelope(
    action: str,
    success: bool = True,
    message: str = "",
    data: dict[str, Any] | None = None,
    error: str | None = None,
) -> dict[str, Any]:
    return {
        "action": action,
        "success": success,
        "message": message,
        "data": data or {},
        "error": error,
    }


def existing_pptx(path_str: str) -> Path:
    p = Path(path_str)
    if not p.exists():
        raise argparse.ArgumentTypeError(f"Datei nicht gefunden: {path_str}")
    if p.suffix.lower() != ".pptx":
        raise argparse.ArgumentTypeError(f"Keine .pptx-Datei: {path_str}")
    return p


def get_shape_type_name(shape: Any) -> str:
    if shape.is_placeholder:
        return "placeholder"
    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.TEXT_BOX:
        return "text_box"
    if st == MSO_SHAPE_TYPE.TABLE:
        return "table"
    if st == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
        return "auto_shape"
    if st == MSO_SHAPE_TYPE.GROUP:
        return "group"
    return "other"


def get_placeholder_type_name(shape: Any) -> str | None:
    if not shape.is_placeholder:
        return None
    try:
        pt = shape.placeholder_format.type
        return str(pt).split(".")[-1].split(" ")[0]
    except Exception:
        return "UNKNOWN"


def cmd_text(infile: Path, include_notes: bool = True, as_json: bool = False) -> None:
    prs = Presentation(str(infile))
    slides_data: list[dict[str, Any]] = []
    text_lines: list[str] = []

    for idx, slide in enumerate(prs.slides, start=1):
        slide_title = ""
        if slide.shapes.title and slide.shapes.title.text:
            slide_title = slide.shapes.title.text.strip()

        shapes_text: list[str] = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        indent = "  " * p.level
                        shapes_text.append(f"{indent}- {t}" if p.level > 0 else t)
            elif shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        shapes_text.append(" | ".join(cells))

        notes_text = ""
        if include_notes and slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()

        slides_data.append({
            "slide_number": idx,
            "title": slide_title,
            "content": shapes_text,
            "speaker_notes": notes_text if include_notes else None,
        })

        title_display = f": {slide_title}" if slide_title else ""
        text_lines.append(f"--- Slide {idx}{title_display} ---")
        if slide_title:
            text_lines.append(f"# {slide_title}")
        for st in shapes_text:
            text_lines.append(st)
        if notes_text:
            text_lines.append(f"> Notes: {notes_text}")
        text_lines.append("")

    if as_json:
        env = make_envelope(
            action="text",
            success=True,
            message=f"Extracted text from {len(slides_data)} slides",
            data={"file": str(infile).replace("\\", "/"), "slides": slides_data},
        )
        print(json.dumps(env, ensure_ascii=False, indent=2))
    else:
        print("\n".join(text_lines).strip())


def cmd_stats(infile: Path, as_json: bool = False) -> None:
    prs = Presentation(str(infile))
    slide_count = len(prs.slides)
    total_words = 0
    shape_counts = {
        "text_boxes": 0,
        "tables": 0,
        "images": 0,
        "placeholders": 0,
        "auto_shapes": 0,
        "groups": 0,
        "other": 0,
    }
    empty_placeholders = 0
    slides_with_notes = 0

    for slide in prs.slides:
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            slides_with_notes += 1
            total_words += len(slide.notes_slide.notes_text_frame.text.split())

        for shape in slide.shapes:
            stype = get_shape_type_name(shape)
            if shape.is_placeholder:
                shape_counts["placeholders"] += 1
                if shape.has_text_frame and not shape.text_frame.text.strip():
                    empty_placeholders += 1
            elif stype == "text_box":
                shape_counts["text_boxes"] += 1
            elif stype == "table":
                shape_counts["tables"] += 1
            elif stype == "picture":
                shape_counts["images"] += 1
            elif stype == "auto_shape":
                shape_counts["auto_shapes"] += 1
            elif stype == "group":
                shape_counts["groups"] += 1
            else:
                shape_counts["other"] += 1

            if shape.has_text_frame:
                total_words += len(shape.text_frame.text.split())
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        total_words += len(cell.text.split())

    data = {
        "file": str(infile).replace("\\", "/"),
        "slide_count": slide_count,
        "total_words": total_words,
        "shapes": shape_counts,
        "empty_placeholders": empty_placeholders,
        "slides_with_notes": slides_with_notes,
    }

    if as_json:
        env = make_envelope(
            action="stats",
            success=True,
            message=f"Statistics for {slide_count} slides",
            data=data,
        )
        print(json.dumps(env, ensure_ascii=False, indent=2))
    else:
        print(f"File: {infile}")
        print(f"Slides: {slide_count}")
        print(f"Total words: {total_words}")
        print(f"Slides with notes: {slides_with_notes}")
        print("Shapes breakdown:")
        for k, v in shape_counts.items():
            print(f"  - {k}: {v}")
        print(f"Empty placeholders: {empty_placeholders}")


def cmd_replace(
    infile: Path,
    outfile: Path,
    find_str: str,
    replace_str: str,
    target_slide: int | None = None,
    include_notes: bool = False,
    as_json: bool = False,
) -> None:
    if not find_str:
        raise ValueError("Parameter '--find' darf nicht leer sein.")

    prs = Presentation(str(infile))
    total_replaced = 0
    affected_slides: list[int] = []

    for idx, slide in enumerate(prs.slides, start=1):
        if target_slide is not None and idx != target_slide:
            continue

        slide_changed = False

        # 1. Shapes text frames
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        if find_str in run.text:
                            count = run.text.count(find_str)
                            run.text = run.text.replace(find_str, replace_str)
                            total_replaced += count
                            slide_changed = True
                    # Fallback if find_str spans across runs in paragraph
                    if find_str in p.text and not any(find_str in r.text for r in p.runs):
                        p.text = p.text.replace(find_str, replace_str)
                        total_replaced += 1
                        slide_changed = True

            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                if find_str in run.text:
                                    count = run.text.count(find_str)
                                    run.text = run.text.replace(find_str, replace_str)
                                    total_replaced += count
                                    slide_changed = True
                            if find_str in p.text and not any(find_str in r.text for r in p.runs):
                                p.text = p.text.replace(find_str, replace_str)
                                total_replaced += 1
                                slide_changed = True

        # 2. Speaker notes
        if include_notes and slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            for p in notes_tf.paragraphs:
                for run in p.runs:
                    if find_str in run.text:
                        count = run.text.count(find_str)
                        run.text = run.text.replace(find_str, replace_str)
                        total_replaced += count
                        slide_changed = True
                if find_str in p.text and not any(find_str in r.text for r in p.runs):
                    p.text = p.text.replace(find_str, replace_str)
                    total_replaced += 1
                    slide_changed = True

        if slide_changed:
            affected_slides.append(idx)

    # Atomic write to outfile
    outfile.parent.mkdir(parents=True, exist_ok=True)
    temp_file = tempfile.NamedTemporaryFile(
        dir=str(outfile.parent),
        prefix=f".tmp_{outfile.stem}_",
        suffix=".pptx",
        delete=False,
    )
    temp_path = Path(temp_file.name)
    temp_file.close()

    try:
        prs.save(str(temp_path))
        os.replace(str(temp_path), str(outfile))
    finally:
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass

    data = {
        "in": str(infile).replace("\\", "/"),
        "out": str(outfile).replace("\\", "/"),
        "find": find_str,
        "replace": replace_str,
        "replaced_occurrences": total_replaced,
        "affected_slides": affected_slides,
    }

    if as_json:
        env = make_envelope(
            action="replace",
            success=True,
            message=f"Replaced {total_replaced} occurrence(s) across {len(affected_slides)} slide(s)",
            data=data,
        )
        print(json.dumps(env, ensure_ascii=False, indent=2))
    else:
        print(f"Replaced {total_replaced} occurrence(s) across {len(affected_slides)} slide(s). Output: {outfile}")


def cmd_inspect(infile: Path, target_slide: int | None = None, as_json: bool = False) -> None:
    prs = Presentation(str(infile))
    width_in = prs.slide_width / 914400.0
    height_in = prs.slide_height / 914400.0

    slides_info: list[dict[str, Any]] = []

    for idx, slide in enumerate(prs.slides, start=1):
        if target_slide is not None and idx != target_slide:
            continue

        shapes_info: list[dict[str, Any]] = []
        for shape in slide.shapes:
            stype = get_shape_type_name(shape)
            ph_type = get_placeholder_type_name(shape)

            shape_dict: dict[str, Any] = {
                "shape_id": shape.shape_id,
                "name": shape.name,
                "type": stype,
                "is_placeholder": shape.is_placeholder,
                "placeholder_type": ph_type,
                "position_pt": {
                    "left": round(shape.left / 12700.0, 1),
                    "top": round(shape.top / 12700.0, 1),
                    "width": round(shape.width / 12700.0, 1),
                    "height": round(shape.height / 12700.0, 1),
                },
            }

            if shape.has_text_frame:
                shape_dict["paragraphs_count"] = len(shape.text_frame.paragraphs)
                shape_dict["preview_text"] = shape.text_frame.text[:80].strip()
            elif shape.has_table:
                shape_dict["table"] = {
                    "rows": len(shape.table.rows),
                    "columns": len(shape.table.columns),
                }

            shapes_info.append(shape_dict)

        slide_dict: dict[str, Any] = {
            "slide_number": idx,
            "slide_id": slide.slide_id,
            "layout_name": slide.slide_layout.name,
            "title": slide.shapes.title.text.strip() if (slide.shapes.title and slide.shapes.title.text) else None,
            "shapes_count": len(slide.shapes),
            "has_notes": slide.has_notes_slide and bool(slide.notes_slide.notes_text_frame.text.strip()),
            "shapes": shapes_info,
        }
        slides_info.append(slide_dict)

    result_data = {
        "file": str(infile).replace("\\", "/"),
        "dimensions": {
            "width_inches": round(width_in, 2),
            "height_inches": round(height_in, 2),
        },
        "total_slides": len(prs.slides),
        "inspected_slides": slides_info,
    }

    if as_json:
        env = make_envelope(
            action="inspect",
            success=True,
            message=f"Inspected {len(slides_info)} slide(s)",
            data=result_data,
        )
        print(json.dumps(env, ensure_ascii=False, indent=2))
    else:
        print(f"Presentation: {infile}")
        print(f"Dimensions: {width_in:.2f} x {height_in:.2f} inches ({len(prs.slides)} slides)")
        for s in slides_info:
            print(f"\nSlide {s['slide_number']} (ID: {s['slide_id']}, Layout: '{s['layout_name']}')")
            if s['title']:
                print(f"  Title: {s['title']}")
            print(f"  Shapes ({s['shapes_count']}):")
            for sh in s["shapes"]:
                ph_str = f" [{sh['placeholder_type']}]" if sh['is_placeholder'] else ""
                prev = f" -> '{sh['preview_text']}'" if "preview_text" in sh and sh['preview_text'] else ""
                print(f"    - ID {sh['shape_id']}: {sh['name']} ({sh['type']}{ph_str}){prev}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="PowerPoint (.pptx) CLI Operations")
    sub = parser.add_subparsers(dest="cmd", required=True)

    # text
    p_text = sub.add_parser("text", help="Text aus allen Folien extrahieren")
    p_text.add_argument("--in", dest="infile", type=existing_pptx, required=True)
    p_text.add_argument("--no-notes", action="store_true", help="Vortragsnotizen ausschließen")
    p_text.add_argument("--json", action="store_true", help="Ausgabe als JSON Envelope")

    # stats
    p_stats = sub.add_parser("stats", help="Statistiken der Präsentation berechnen")
    p_stats.add_argument("--in", dest="infile", type=existing_pptx, required=True)
    p_stats.add_argument("--json", action="store_true", help="Ausgabe als JSON Envelope")

    # replace
    p_replace = sub.add_parser("replace", help="Textstellen suchen und ersetzen")
    p_replace.add_argument("--in", dest="infile", type=existing_pptx, required=True)
    p_replace.add_argument("--out", dest="outfile", type=Path, required=True)
    p_replace.add_argument("--find", required=True, help="Suchbegriff")
    p_replace.add_argument("--replace", required=True, help="Ersatztext")
    p_replace.add_argument("--slide", type=int, help="Optional auf bestimmte Folie (1-basiert) einschränken")
    p_replace.add_argument("--include-notes", action="store_true", help="Auch in Vortragsnotizen ersetzen")
    p_replace.add_argument("--json", action="store_true", help="Ausgabe als JSON Envelope")

    # inspect
    p_inspect = sub.add_parser("inspect", help="Detaillierte Folien- und Shape-Inspektion")
    p_inspect.add_argument("--in", dest="infile", type=existing_pptx, required=True)
    p_inspect.add_argument("--slide", type=int, help="Optional auf bestimmte Folie (1-basiert) einschränken")
    p_inspect.add_argument("--json", action="store_true", help="Ausgabe als JSON Envelope")

    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.cmd == "text":
        cmd_text(args.infile, include_notes=not args.no_notes, as_json=args.json)
    elif args.cmd == "stats":
        cmd_stats(args.infile, as_json=args.json)
    elif args.cmd == "replace":
        cmd_replace(
            args.infile,
            args.outfile,
            args.find,
            args.replace,
            target_slide=args.slide,
            include_notes=args.include_notes,
            as_json=args.json,
        )
    elif args.cmd == "inspect":
        cmd_inspect(args.infile, target_slide=args.slide, as_json=args.json)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # noqa: BLE001
        err_env = make_envelope(
            action="error",
            success=False,
            message=str(exc),
            error=str(exc),
        )
        if "--json" in sys.argv:
            print(json.dumps(err_env, ensure_ascii=False, indent=2))
        else:
            print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)
