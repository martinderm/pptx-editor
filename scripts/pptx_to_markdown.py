#!/usr/bin/env python3
"""Convert a PowerPoint (.pptx) presentation into structured Markdown.

Extracts slide titles, bullet hierarchies, tables, and speaker notes
into a clean, human- and LLM-readable Markdown document.

Usage:
  python scripts/pptx_to_markdown.py --in presentation.pptx --out presentation.md
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

try:
    from pptx_ops import load_presentation
except ImportError:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from pptx_ops import load_presentation


def existing_pptx(path_str: str) -> Path:
    p = Path(path_str)
    if not p.exists():
        raise argparse.ArgumentTypeError(f"Datei nicht gefunden: {path_str}")
    if p.suffix.lower() not in (".pptx", ".potx"):
        raise argparse.ArgumentTypeError(f"Keine .pptx- oder .potx-Datei: {path_str}")
    return p


def pptx_to_markdown(prs: Presentation) -> str:
    md_slides: list[str] = []
    footer_text = ""

    for idx, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []

        # Check for footer metadata on first slide
        if idx == 1:
            for shape in slide.shapes:
                if shape.is_placeholder and getattr(shape, "placeholder_format", None):
                    if shape.placeholder_format.type == PP_PLACEHOLDER.FOOTER and shape.text.strip():
                        footer_text = shape.text.strip()

        # 1. Slide Title
        title_text = ""
        if slide.shapes.title and slide.shapes.title.text:
            title_text = slide.shapes.title.text.strip()
            lines.append(f"# {title_text}\n")

        # 2. Content shapes (Textboxes, Body Placeholders)
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue

            # Skip footer, slide number and date placeholders from bullet extraction
            if shape.is_placeholder and getattr(shape, "placeholder_format", None):
                if shape.placeholder_format.type in (
                    PP_PLACEHOLDER.FOOTER,
                    PP_PLACEHOLDER.SLIDE_NUMBER,
                    PP_PLACEHOLDER.DATE,
                ):
                    continue

            if shape.has_table:
                tbl = shape.table
                rows_data: list[list[str]] = []
                for row in tbl.rows:
                    row_cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
                    rows_data.append(row_cells)

                if rows_data:
                    header = rows_data[0]
                    lines.append("| " + " | ".join(header) + " |")
                    lines.append("| " + " | ".join(["---"] * len(header)) + " |")
                    for r in rows_data[1:]:
                        lines.append("| " + " | ".join(r) + " |")
                    lines.append("")

            elif shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if not t:
                        continue
                    indent = "  " * p.level
                    lines.append(f"{indent}- {t}")

        # 3. Speaker Notes
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            notes_text = notes_tf.text.strip()
            if notes_text:
                lines.append("")
                lines.append(f"> Notes: {notes_text}")

        md_slides.append("\n".join(lines).strip())

    body = "\n\n---\n\n".join(md_slides) + "\n"
    if footer_text:
        parts = [p.strip() for p in footer_text.split("|")]
        if len(parts) == 3:
            fm = f"---\npresenter: \"{parts[0]}\"\nevent: \"{parts[1]}\"\ndate: \"{parts[2]}\"\n---\n\n"
        else:
            fm = f"---\nfooter: \"{footer_text}\"\n---\n\n"
        return fm + body
    return body


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert PowerPoint (.pptx) to Markdown")
    parser.add_argument("--in", dest="infile", type=existing_pptx, required=True, help="Eingabe PPTX-Datei")
    parser.add_argument("--out", dest="outfile", type=Path, required=True, help="Ausgabe Markdown-Datei")
    args = parser.parse_args()

    prs = load_presentation(args.infile)
    md_content = pptx_to_markdown(prs)

    # Atomic write
    args.outfile.parent.mkdir(parents=True, exist_ok=True)
    temp_file = tempfile.NamedTemporaryFile(
        dir=str(args.outfile.parent),
        prefix=f".tmp_{args.outfile.stem}_",
        suffix=".md",
        delete=False,
    )
    temp_path = Path(temp_file.name)
    temp_file.close()

    try:
        temp_path.write_text(md_content, encoding="utf-8", newline="\n")
        os.replace(str(temp_path), str(args.outfile))
    finally:
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass

    out_info = {
        "action": "pptx_to_markdown",
        "success": True,
        "message": f"Exported {len(prs.slides)} slides to Markdown",
        "in": str(args.infile).replace("\\", "/"),
        "out": str(args.outfile).replace("\\", "/"),
        "slides_count": len(prs.slides),
    }
    print(json.dumps(out_info, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # noqa: BLE001
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)
